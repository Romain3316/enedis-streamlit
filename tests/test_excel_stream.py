"""Compare streamed workbook values and formatting to the original pandas export."""
from datetime import date
from copy import copy
from io import BytesIO

import pandas as pd
import pytest
from openpyxl import load_workbook
from excel_stream import report_workbook


@pytest.mark.parametrize('include_index', [False, True])
def test_streamed_export_matches_pandas(include_index):
    frame = pd.DataFrame({
        'Date': pd.to_datetime(['2024-09-03 01:05', None, '2024-10-27 02:00']),
        'Énergie': [0.0, 12.345, float('nan')],
        'Tarif': [float('inf'), float('-inf'), 0.1234],
        'Texte': ['électricité', '', 'HP'],
        'Jour': [date(2024, 9, 3), None, date(2024, 10, 27)],
    }, index=pd.Index(['Lundi', 'Mardi', 'Mercredi'], name='Jour semaine'))
    expected = BytesIO()
    with pd.ExcelWriter(expected, engine='openpyxl') as writer:
        frame.to_excel(writer, sheet_name='Synthèse', index=include_index)
        sheet = writer.book['Synthèse']
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            width = max(len(str(cell.value or '')) for cell in column)
            sheet.column_dimensions[column[0].column_letter].width = min(max(width + 2, 11), 35)
    old = load_workbook(BytesIO(expected.getvalue()))['Synthèse']
    new = load_workbook(BytesIO(report_workbook([('Synthèse', frame, include_index)])))['Synthèse']
    assert old.max_row == new.max_row
    assert old.max_column == new.max_column
    assert old.freeze_panes == new.freeze_panes
    assert old.auto_filter.ref == new.auto_filter.ref
    for old_row, new_row in zip(old, new):
        for a, b in zip(old_row, new_row):
            assert a.value == b.value
            assert a.number_format == b.number_format
            assert copy(a.font) == copy(b.font)
            assert copy(a.border) == copy(b.border)
            assert copy(a.alignment) == copy(b.alignment)
    for column in old.column_dimensions:
        assert old.column_dimensions[column].width == new.column_dimensions[column].width
