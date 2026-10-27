"""Write the report workbook one row at a time, without a full cell graph."""
from datetime import date, datetime
from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter


def report_workbook(sheets):
    output = BytesIO()
    workbook = Workbook(write_only=True)
    header_font = Font(bold=True)
    header_border = Border(**{edge: Side(style="thin") for edge in ("top", "right", "bottom", "left")})
    header_alignment = Alignment(horizontal="center", vertical="top")

    def value_for_excel(value):
        if pd.isna(value):
            return ""
        if isinstance(value, float) and value in (float("inf"), float("-inf")):
            return str(value)
        return value

    for name, frame, include_index in sheets:
        sheet = workbook.create_sheet(name)
        headers = ([frame.index.name] if include_index else []) + list(frame.columns)

        def rows():
            yield headers
            for row in frame.itertuples(index=include_index, name=None):
                yield tuple(value_for_excel(value) for value in row)

        # Column widths must precede the streamed rows in the XLSX XML.
        widths = [0] * len(headers)
        for row in rows():
            for column, value in enumerate(row):
                widths[column] = max(widths[column], len(str(value or "")))
        for column, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(column)].width = min(max(width + 2, 11), 35)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(max(len(headers), 1))}{len(frame) + 1}"
        for row_number, row in enumerate(rows()):
            cells = []
            for column, value in enumerate(row):
                cell = WriteOnlyCell(sheet, value=value)
                if row_number == 0 or (include_index and column == 0):
                    if value is not None:
                        cell.font = header_font
                        cell.border = header_border
                        cell.alignment = header_alignment
                if isinstance(value, datetime):
                    cell.number_format = "YYYY-MM-DD HH:MM:SS"
                elif isinstance(value, date):
                    cell.number_format = "YYYY-MM-DD"
                cells.append(cell)
            sheet.append(cells)
    workbook.save(output)
    return output.getvalue()
